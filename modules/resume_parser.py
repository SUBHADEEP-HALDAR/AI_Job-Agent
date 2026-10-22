"""Resume text extraction and rule-based profile building.

Design notes
- Everything runs locally; resume text is never logged or sent anywhere.
- The parser is deliberately simple and rule-based. The user always reviews and
  corrects the profile before matching, so extraction mistakes are expected.
"""
import re
from io import BytesIO

from . import skills as sk

MAX_BYTES = 5 * 1024 * 1024  # 5 MB
MIN_TEXT_CHARS = 150  # below this a PDF is probably scanned


class ResumeError(Exception):
    """Raised with a user-friendly message when a resume can't be read."""


# ---------------------------------------------------------------- extraction
def extract_text(file_bytes: bytes, filename: str) -> str:
    if not file_bytes:
        raise ResumeError("The file is empty.")
    if len(file_bytes) > MAX_BYTES:
        raise ResumeError("The file is larger than 5 MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        if not file_bytes.startswith(b"%PDF"):
            raise ResumeError("This file is not a valid PDF.")
        return _pdf_text(file_bytes)
    if name.endswith(".docx"):
        if not file_bytes.startswith(b"PK"):
            raise ResumeError("This file is not a valid DOCX.")
        return _docx_text(file_bytes)
    raise ResumeError("Only PDF and DOCX files are supported.")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ResumeError("This PDF is password-protected.")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except ResumeError:
        raise
    except Exception:
        raise ResumeError("Could not read this PDF (it may be damaged).")
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ResumeError(
            "Very little text could be read. This looks like a scanned PDF "
            "(an image). Export a text-based PDF/DOCX from your editor, or use "
            "manual entry below."
        )
    return text


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(BytesIO(data))
    except Exception:
        raise ResumeError("Could not read this DOCX (it may be damaged).")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)
    if len(text.strip()) < 50:
        raise ResumeError("Very little text was found in this DOCX.")
    return text


# ------------------------------------------------------------------ sections
_SECTION_PATTERNS = {
    "skills": r"(technical\s+)?skills?|core competenc(y|ies)|technologies|tech stack|skills? summary",
    "education": r"education(al)?( background| qualifications?)?|academics?|qualifications?",
    "projects": r"(academic |personal |key )?projects?",
    "internships": r"internships?|training",
    "experience": r"(work |professional |employment )?experience|employment( history)?",
    "certifications": r"certifications?|certificates?|courses|licen[sc]es",
    "other": r"summary|objective|career objective|profile|about me|achievements?|awards?|"
             r"interests?|hobbies|references?|declaration|languages?|extracurriculars?|activities",
}
_SECTION_RES = {k: re.compile(v, re.I) for k, v in _SECTION_PATTERNS.items()}
_BULLETS = "•●■◦▪‣·-–—*>"


def _heading_of(line: str):
    t = re.sub(r"[^A-Za-z &/]", "", line).strip()
    if not t or len(line.strip()) > 40:
        return None
    for key, rx in _SECTION_RES.items():
        if rx.fullmatch(t):
            return key
    return None


def split_sections(text: str) -> dict:
    sections = {k: [] for k in _SECTION_PATTERNS}
    current = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        h = _heading_of(line)
        if h:
            current = h
            continue
        if current:
            item = line.lstrip(_BULLETS + " \t").strip()
            if len(item) >= 3:
                sections[current].append(item)
    return sections


# ----------------------------------------------------------------- education
_EDU_CI = [
    (4, re.compile(r"ph\.?\s?d|doctorate", re.I)),
    (3, re.compile(r"m\.?\s?tech|m\.?\s?sc|\bmca\b|\bmba\b|master|\bm\.?\s?com\b", re.I)),
    (2, re.compile(r"b\.?\s?tech|b\.?\s?sc|\bbca\b|\bbba\b|bachelor|\bb\.?\s?com\b|undergraduate", re.I)),
    (1, re.compile(r"diploma|polytechnic", re.I)),
]
# ambiguous abbreviations are matched case-sensitively (so "be" / "ma" don't match)
_EDU_CS = [
    (3, re.compile(r"\bM\.?E\b\.?|\bM\.?A\b\.?")),
    (2, re.compile(r"\bB\.?E\b\.?|\bB\.?A\b\.?")),
]


def education_level(text: str):
    """Highest degree level found (0-4) or None if nothing recognisable.

    1 diploma, 2 bachelor's, 3 master's, 4 PhD.
    """
    found = []
    for level, rx in _EDU_CI:
        if rx.search(text or ""):
            found.append(level)
    for level, rx in _EDU_CS:
        if rx.search(text or ""):
            found.append(level)
    return max(found) if found else None


EDU_LABELS = {None: "not found", 0: "none", 1: "diploma", 2: "bachelor's", 3: "master's", 4: "PhD"}

# ---------------------------------------------------------------- experience
_YEARS_RX = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:[\w/&-]+\s+){0,3}?experience",
    re.I,
)


def guess_years_experience(text: str) -> float:
    vals = [float(m.group(1)) for m in _YEARS_RX.finditer(text or "")]
    vals = [v for v in vals if v <= 40]
    return max(vals) if vals else 0.0


# ------------------------------------------------------------------- profile
def empty_profile() -> dict:
    return {
        "skills": [],
        "education": [],
        "projects": [],
        "internships": [],
        "experience": [],
        "certifications": [],
        "years_experience": 0.0,
    }


def _skills_section_extras(lines: list, known: set) -> list:
    extras = []
    for line in lines:
        if ":" in line:
            line = line.split(":", 1)[1]
        for tok in re.split(r"[,;|/•·]", line):
            tok = tok.strip(" .-\t")
            if not (2 <= len(tok) <= 25):
                continue
            if tok.lower() in known or sk.canon(tok).lower() in known or tok in extras:
                continue
            extras.append(tok)
    return extras[:30]


def parse_resume_text(text: str) -> dict:
    """Turn resume text into a candidate profile dict (to be reviewed by the user)."""
    sections = split_sections(text)
    profile = empty_profile()

    found = sk.find_skills(text)
    known = {s.lower() for s in found}
    extras = _skills_section_extras(sections["skills"], known)
    profile["skills"] = found + extras

    edu = sections["education"][:8]
    if not edu:
        edu = [ln.strip() for ln in text.splitlines()
               if ln.strip() and education_level(ln) is not None][:6]
    profile["education"] = edu

    for key in ("projects", "internships", "experience", "certifications"):
        profile[key] = sections[key][:40]

    profile["years_experience"] = guess_years_experience(text)
    return profile


def profile_evidence_text(profile: dict) -> str:
    parts = []
    for key in ("projects", "internships", "experience", "certifications"):
        parts.extend(profile.get(key, []))
    return " ".join(parts).lower()
